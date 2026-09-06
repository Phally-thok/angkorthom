import os
import re
import sqlite3
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file, jsonify

from database import get_db_connection, init_db, generate_password_hash, check_password_hash

app = Flask(__name__)
app.secret_key = 'angkor_thom_education_secret_key_2026'
EXPORTS_DIR = os.path.join(os.path.dirname(__file__), 'exports')
os.makedirs(EXPORTS_DIR, exist_ok=True)

# Ensure DB is initialized on start
init_db()

# Login Decorator
def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            flash('សូមចូលប្រើប្រាស់ប្រព័ន្ធជាមុនសិន!', 'danger')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

# Role Check Decorator
def role_required(*roles):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if 'role' not in session or session['role'] not in roles:
                flash('លោកអ្នកគ្មានសិទ្ធិចូលប្រើប្រាស់មុខងារនេះទេ!', 'danger')
                return redirect(url_for('dashboard'))
            return f(*args, **kwargs)
        return decorated_function
    return decorator

# --- AUTH ROUTES ---
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')

        conn = get_db_connection()
        user = conn.execute('SELECT * FROM users WHERE username = ?', (username,)).fetchone()
        conn.close()

        if user and check_password_hash(user['password_hash'], password):
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['user_name'] = user['full_name']
            session['role'] = user['role']
            session['school_id'] = user['school_id']
            flash(f'ស្វាគមន៍ {user["full_name"]} មកកាន់ប្រព័ន្ធគ្រប់គ្រង!', 'success')
            return redirect(url_for('dashboard'))
        else:
            flash('ឈ្មោះគណនី ឬ ពាក្យសម្ងាត់មិនត្រឹមត្រូវទេ!', 'danger')

    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    flash('លោកអ្នកបានចាកចេញពីប្រព័ន្ធដោយជោគជ័យ!', 'success')
    return redirect(url_for('login'))

# --- DASHBOARD ---
@app.route('/')
@app.route('/dashboard')
@login_required
def dashboard():
    conn = get_db_connection()

    total_schools = conn.execute('SELECT COUNT(*) FROM schools').fetchone()[0]
    total_staff = conn.execute('SELECT COUNT(*) FROM staff WHERE status = "active"').fetchone()[0]
    staff_a = conn.execute('SELECT COUNT(*) FROM staff WHERE framework = "A"').fetchone()[0]
    staff_b = conn.execute('SELECT COUNT(*) FROM staff WHERE framework = "B"').fetchone()[0]
    staff_c = conn.execute('SELECT COUNT(*) FROM staff WHERE framework = "C"').fetchone()[0]

    # Fetch distinct academic years ordered DESCENDING (newest year at top)
    years_rows = conn.execute('SELECT DISTINCT academic_year FROM student_stats ORDER BY academic_year DESC').fetchall()
    academic_years = [r[0] for r in years_rows if r[0]]
    latest_year = academic_years[0] if academic_years else '2025-2026'

    student_sum = conn.execute('''
        SELECT SUM(total_students), SUM(female_students) FROM student_stats WHERE academic_year = ?
    ''', (latest_year,)).fetchone()
    total_students = student_sum[0] or 0
    female_students = student_sum[1] or 0

    exam_bacii = conn.execute('''
        SELECT SUM(total_candidates), SUM(total_passed) FROM exam_results WHERE exam_type = "bacii_grade_12"
    ''').fetchone()
    bacii_cand = exam_bacii[0] or 1
    bacii_passed = exam_bacii[1] or 0
    bacii_pass_rate = round((bacii_passed / bacii_cand) * 100, 1)

    # Function to get detailed per-school per-academic-year chart data
    def get_detailed_level_data(level_name):
        schools_in_level = conn.execute('SELECT id, name_kh FROM schools WHERE level = ? ORDER BY id ASC', (level_name,)).fetchall()
        rows = conn.execute('''
            SELECT s.name_kh, st.academic_year, st.total_students, st.female_students
            FROM student_stats st
            JOIN schools s ON st.school_id = s.id
            WHERE s.level = ?
            ORDER BY s.id ASC, st.academic_year ASC
        ''', (level_name,)).fetchall()

        by_year = {}
        for yr in academic_years:
            by_year[yr] = []
            for sch in schools_in_level:
                s_name = sch['name_kh']
                match = next((r for r in rows if r['name_kh'] == s_name and r['academic_year'] == yr), None)
                by_year[yr].append({
                    'school_name': s_name,
                    'total': match['total_students'] if match else 0,
                    'female': match['female_students'] if match else 0
                })

        all_years_flat = {'labels': [], 'totals': [], 'females': []}
        for sch in schools_in_level:
            s_name = sch['name_kh']
            short_name = s_name.replace('សាលាបឋមសិក្សា ', '').replace('អនុវិទ្យាល័យ ', '').replace('វិទ្យាល័យ ', '').replace('សាលាមត្តេយ្យសិក្សារដ្ឋ ', '').replace('សាលាមត្តេយ្យសិក្សាសហគមន៍ ', '').replace('សាលាមត្តេយ្យសិក្សា ', '')
            for yr in academic_years:
                match = next((r for r in rows if r['name_kh'] == s_name and r['academic_year'] == yr), None)
                all_years_flat['labels'].append(f'{short_name} ({yr})')
                all_years_flat['totals'].append(match['total_students'] if match else 0)
                all_years_flat['females'].append(match['female_students'] if match else 0)

        return {
            'schools': [s['name_kh'] for s in schools_in_level],
            'years': academic_years,
            'by_year': by_year,
            'all_years_flat': all_years_flat
        }

    community_preschool_chart = get_detailed_level_data('community_preschool')
    state_preschool_chart = get_detailed_level_data('state_preschool')
    primary_chart = get_detailed_level_data('primary')
    secondary_chart = get_detailed_level_data('secondary')
    high_chart = get_detailed_level_data('high')

    # School level counts summary
    school_counts = {
        'community_preschool': conn.execute("SELECT COUNT(*) FROM schools WHERE level = 'community_preschool'").fetchone()[0],
        'state_preschool': conn.execute("SELECT COUNT(*) FROM schools WHERE level IN ('state_preschool', 'preschool')").fetchone()[0],
        'primary': conn.execute("SELECT COUNT(*) FROM schools WHERE level = 'primary'").fetchone()[0],
        'secondary': conn.execute("SELECT COUNT(*) FROM schools WHERE level = 'secondary'").fetchone()[0],
        'high': conn.execute("SELECT COUNT(*) FROM schools WHERE level = 'high'").fetchone()[0],
    }

    # Exam Chart Data by Academic Year
    exam_rows = conn.execute('''
        SELECT academic_year,
               SUM(grade_a) as a, SUM(grade_b) as b, SUM(grade_c) as c, SUM(grade_d) as d, SUM(grade_e) as e
        FROM exam_results
        GROUP BY academic_year
    ''').fetchall()

    exam_chart_by_year = {}
    for yr in academic_years:
        match = next((r for r in exam_rows if r['academic_year'] == yr), None)
        if match and (match['a'] or match['b'] or match['c'] or match['d'] or match['e']):
            exam_chart_by_year[yr] = [match['a'] or 0, match['b'] or 0, match['c'] or 0, match['d'] or 0, match['e'] or 0]
        else:
            exam_chart_by_year[yr] = [0, 0, 0, 0, 0]

    conn.close()

    stats = {
        'total_schools': total_schools,
        'total_staff': total_staff,
        'staff_framework_a': staff_a,
        'staff_framework_b': staff_b,
        'staff_framework_c': staff_c,
        'total_students': total_students,
        'female_students': female_students,
        'bacii_pass_rate': bacii_pass_rate,
        'school_counts': school_counts,
        'latest_year': latest_year
    }

    return render_template(
        'dashboard.html',
        active_page='dashboard',
        stats=stats,
        community_preschool_chart=community_preschool_chart,
        state_preschool_chart=state_preschool_chart,
        primary_chart=primary_chart,
        secondary_chart=secondary_chart,
        high_chart=high_chart,
        academic_years=academic_years,
        exam_chart_by_year=exam_chart_by_year
    )

@app.route('/academic_year/add', methods=['POST'])
@login_required
@role_required('admin', 'user')
def add_academic_year():
    academic_year = request.form.get('academic_year')
    copy_previous = request.form.get('copy_previous', '1')

    if not academic_year:
        flash('សូមបញ្ចូលឆ្នាំសិក្សា!', 'danger')
        return redirect(url_for('dashboard'))

    academic_year = academic_year.strip()
    conn = get_db_connection()

    existing = conn.execute('SELECT COUNT(*) FROM student_stats WHERE academic_year = ?', (academic_year,)).fetchone()[0]
    if existing > 0:
        flash(f'ឆ្នាំសិក្សា {academic_year} មាននៅក្នុងប្រព័ន្ធរួចហើយ!', 'warning')
        conn.close()
        return redirect(url_for('dashboard'))

    schools = conn.execute('SELECT id FROM schools').fetchall()
    for sch in schools:
        school_id = sch['id']
        prev_stat = None
        if copy_previous == '1':
            prev_stat = conn.execute('SELECT * FROM student_stats WHERE school_id = ? ORDER BY id DESC LIMIT 1', (school_id,)).fetchone()
        
        if prev_stat:
            conn.execute('''
                INSERT INTO student_stats (school_id, academic_year, total_students, female_students, disabled_students, disabled_female, scholarship_students, scholarship_female, dropout_count, dropout_female, repetition_count)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (school_id, academic_year, prev_stat['total_students'], prev_stat['female_students'], prev_stat['disabled_students'], prev_stat['disabled_female'], prev_stat['scholarship_students'], prev_stat['scholarship_female'], prev_stat['dropout_count'], prev_stat['dropout_female'], prev_stat['repetition_count']))
        else:
            conn.execute('''
                INSERT INTO student_stats (school_id, academic_year, total_students, female_students)
                VALUES (?, ?, 0, 0)
            ''', (school_id, academic_year))

    conn.commit()
    conn.close()
    flash(f'បានបន្ថែមឆ្នាំសិក្សា {academic_year} ចូលក្នុងប្រព័ន្ធដោយជោគជ័យ!', 'success')
    return redirect(url_for('dashboard'))

# --- MODULE 1: SCHOOL MANAGEMENT ---
@app.route('/schools')
@login_required
@role_required('admin', 'user')
def schools():
    action = request.args.get('action')
    level_filter = request.args.get('level', 'all')
    if action == 'delete':
        target = request.args.get('target')
        item_id = request.args.get('id')
        if item_id and target == 'school':
            conn = get_db_connection()
            conn.execute('DELETE FROM schools WHERE id = ?', (item_id,))
            conn.commit()
            conn.close()
            flash('បានលុបទិន្នន័យសាលារៀនដោយជោគជ័យ!', 'success')
            if level_filter and level_filter != 'all':
                return redirect(url_for('schools', level=level_filter))
            return redirect(url_for('schools'))

    conn = get_db_connection()
    schools_list = conn.execute('SELECT * FROM schools ORDER BY level DESC, code ASC').fetchall()
    
    level_counts = {
        'all': len(schools_list),
        'community_preschool': conn.execute("SELECT COUNT(*) FROM schools WHERE level = 'community_preschool'").fetchone()[0],
        'state_preschool': conn.execute("SELECT COUNT(*) FROM schools WHERE level IN ('state_preschool', 'preschool')").fetchone()[0],
        'primary': conn.execute("SELECT COUNT(*) FROM schools WHERE level = 'primary'").fetchone()[0],
        'secondary': conn.execute("SELECT COUNT(*) FROM schools WHERE level = 'secondary'").fetchone()[0],
        'high': conn.execute("SELECT COUNT(*) FROM schools WHERE level = 'high'").fetchone()[0],
    }
    conn.close()
    return render_template('schools.html', active_page='schools', schools=schools_list, level_counts=level_counts, level_filter=level_filter)

@app.route('/schools/add', methods=['POST'])
@login_required
@role_required('admin', 'chief')
def add_school():
    school_id = request.form.get('school_id')
    code = request.form.get('code')
    name_kh = request.form.get('name_kh')
    name_en = request.form.get('name_en')
    level = request.form.get('level')
    commune = request.form.get('commune')
    village = request.form.get('village')
    principal_name = request.form.get('principal_name')
    principal_phone = request.form.get('principal_phone')
    classrooms = request.form.get('classrooms', 0)
    buildings = request.form.get('buildings', 0)

    conn = get_db_connection()
    if school_id:
        conn.execute('''
            UPDATE schools SET code=?, name_kh=?, name_en=?, level=?, commune=?, village=?, principal_name=?, principal_phone=?, classrooms=?, buildings=?
            WHERE id=?
        ''', (code, name_kh, name_en, level, commune, village, principal_name, principal_phone, classrooms, buildings, school_id))
        flash('បានធ្វើបច្ចុប្បន្នភាពព័ត៌មានសាលារៀនដោយជោគជ័យ!', 'success')
    else:
        conn.execute('''
            INSERT INTO schools (code, name_kh, name_en, level, commune, village, principal_name, principal_phone, classrooms, buildings)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (code, name_kh, name_en, level, commune, village, principal_name, principal_phone, classrooms, buildings))
        flash('បានបន្ថែមសាលារៀនថ្មីដោយជោគជ័យ!', 'success')

    conn.commit()
    conn.close()
    return redirect(url_for('schools', level=level))

# --- MODULE 2: STAFF & TEACHER MANAGEMENT ---
@app.route('/staff')
@login_required
def staff():
    action = request.args.get('action')
    if action == 'delete':
        target = request.args.get('target')
        item_id = request.args.get('id')
        if item_id:
            conn = get_db_connection()
            if target == 'staff':
                conn.execute('DELETE FROM staff WHERE id = ?', (item_id,))
                flash('បានលុបទិន្នន័យគ្រូបង្រៀន/បុគ្គលិកសាលាដោយជោគជ័យ!', 'success')
            elif target == 'office_staff':
                conn.execute('DELETE FROM office_staff WHERE id = ?', (item_id,))
                flash('បានលុបទិន្នន័យបុគ្គលិកការិយាល័យដោយជោគជ័យ!', 'success')
            conn.commit()
            conn.close()
        sub = request.args.get('sub')
        if sub:
            return redirect(url_for('staff', sub=sub))
        return redirect(url_for('staff'))

    conn = get_db_connection()
    staff_list = conn.execute('''
        SELECT st.*, sc.name_kh as school_name
        FROM staff st
        JOIN schools sc ON st.school_id = sc.id
        ORDER BY st.id DESC
    ''').fetchall()

    office_staff_list = conn.execute('''
        SELECT * FROM office_staff ORDER BY id DESC
    ''').fetchall()

    history_list = conn.execute('''
        SELECT sh.*, st.name_kh as staff_name
        FROM staff_history sh
        JOIN staff st ON sh.staff_id = st.id
        ORDER BY sh.id DESC
    ''').fetchall()

    leaves_list = conn.execute('''
        SELECT sl.*, st.name_kh as staff_name
        FROM staff_leaves sl
        JOIN staff st ON sl.staff_id = st.id
        ORDER BY sl.id DESC
    ''').fetchall()

    schools_list = conn.execute('SELECT id, name_kh FROM schools').fetchall()
    conn.close()

    return render_template(
        'staff.html',
        active_page='staff',
        staff_list=staff_list,
        office_staff_list=office_staff_list,
        history_list=history_list,
        leaves_list=leaves_list,
        schools=schools_list
    )

@app.route('/staff/add', methods=['POST'])
@login_required
@role_required('admin', 'chief')
def add_staff():
    staff_id = request.form.get('staff_id')
    staff_id_num = request.form.get('staff_id_num')
    name_kh = request.form.get('name_kh')
    name_en = request.form.get('name_en')
    sex = request.form.get('sex')
    school_id = request.form.get('school_id')
    framework = request.form.get('framework')
    position = request.form.get('position')
    subject_specialty = request.form.get('subject_specialty')
    current_grade = request.form.get('current_grade')
    phone = request.form.get('phone')

    conn = get_db_connection()
    if staff_id:
        conn.execute('''
            UPDATE staff SET staff_id_num=?, name_kh=?, name_en=?, sex=?, school_id=?, framework=?, position=?, subject_specialty=?, current_grade=?, phone=?
            WHERE id=?
        ''', (staff_id_num, name_kh, name_en, sex, school_id, framework, position, subject_specialty, current_grade, phone, staff_id))
        flash('បានកែប្រែប្រវត្តិរូបគ្រូបង្រៀន/បុគ្គលិកដោយជោគជ័យ!', 'success')
    else:
        conn.execute('''
            INSERT INTO staff (staff_id_num, name_kh, name_en, sex, school_id, framework, position, subject_specialty, current_grade, phone)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (staff_id_num, name_kh, name_en, sex, school_id, framework, position, subject_specialty, current_grade, phone))
        flash('បានបន្ថែមប្រវត្តិរូបគ្រូបង្រៀន/បុគ្គលិកថ្មីដោយជោគជ័យ!', 'success')
    conn.commit()
    conn.close()
    return redirect(url_for('staff', sub='school'))

@app.route('/office_staff/add', methods=['POST'])
@login_required
@role_required('admin', 'chief')
def add_office_staff():
    office_staff_id = request.form.get('office_staff_id')
    staff_id_num = request.form.get('staff_id_num')
    name_kh = request.form.get('name_kh')
    name_en = request.form.get('name_en')
    sex = request.form.get('sex')
    dob = request.form.get('dob')
    phone = request.form.get('phone')
    email = request.form.get('email')
    office_unit = request.form.get('office_unit')
    framework = request.form.get('framework')
    position = request.form.get('position')
    current_grade = request.form.get('current_grade')
    salary_index = request.form.get('salary_index', 1.0)
    bank_account = request.form.get('bank_account')
    status = request.form.get('status', 'active')

    conn = get_db_connection()
    if office_staff_id:
        conn.execute('''
            UPDATE office_staff SET staff_id_num=?, name_kh=?, name_en=?, sex=?, dob=?, phone=?, email=?, office_unit=?, framework=?, position=?, current_grade=?, salary_index=?, bank_account=?, status=?
            WHERE id=?
        ''', (staff_id_num, name_kh, name_en, sex, dob, phone, email, office_unit, framework, position, current_grade, salary_index, bank_account, status, office_staff_id))
        flash('បានកែប្រែប្រវត្តិរូបបុគ្គលិកការិយាល័យដោយជោគជ័យ!', 'success')
    else:
        conn.execute('''
            INSERT INTO office_staff (staff_id_num, name_kh, name_en, sex, dob, phone, email, office_unit, framework, position, current_grade, salary_index, bank_account, status)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (staff_id_num, name_kh, name_en, sex, dob, phone, email, office_unit, framework, position, current_grade, salary_index, bank_account, status))
        flash('បានបន្ថែមប្រវត្តិរូបបុគ្គលិកការិយាល័យថ្មីដោយជោគជ័យ!', 'success')
    conn.commit()
    conn.close()
    return redirect(url_for('staff', sub='office'))

@app.route('/staff/history/add', methods=['POST'])
@login_required
@role_required('admin', 'chief')
def add_staff_history():
    staff_id = request.form.get('staff_id')
    change_type = request.form.get('change_type')
    effective_date = request.form.get('effective_date')
    old_detail = request.form.get('old_detail')
    new_detail = request.form.get('new_detail')
    reference_doc = request.form.get('reference_doc')

    conn = get_db_connection()
    conn.execute('''
        INSERT INTO staff_history (staff_id, change_type, effective_date, old_detail, new_detail, reference_doc)
        VALUES (?, ?, ?, ?, ?, ?)
    ''', (staff_id, change_type, effective_date, old_detail, new_detail, reference_doc))
    conn.commit()
    conn.close()
    flash('បានកត់ត្រាប្រវត្តិតំឡើងថ្នាក់/ផ្ទេរដោយជោគជ័យ!', 'success')
    return redirect(url_for('staff'))

@app.route('/staff/leave/add', methods=['POST'])
@login_required
def add_staff_leave():
    staff_id = request.form.get('staff_id')
    leave_type = request.form.get('type')
    title = request.form.get('title')
    start_date = request.form.get('start_date')
    end_date = request.form.get('end_date')
    location_or_reason = request.form.get('location_or_reason')

    conn = get_db_connection()
    conn.execute('''
        INSERT INTO staff_leaves (staff_id, type, title, start_date, end_date, location_or_reason, status)
        VALUES (?, ?, ?, ?, ?, ?, 'approved')
    ''', (staff_id, leave_type, title, start_date, end_date, location_or_reason))
    conn.commit()
    conn.close()
    flash('បានកត់ត្រាការសុំច្បាប់/វគ្គបណ្តុះបណ្តាលដោយជោគជ័យ!', 'success')
    return redirect(url_for('staff'))

# --- MODULE 3: ស្ថិតិសិស្ស និងលទ្ធផលប្រឡង ---
@app.route('/students')
@login_required
def students():
    level_filter = request.args.get('level', 'all')
    action = request.args.get('action')
    if action == 'delete':
        target = request.args.get('target')
        item_id = request.args.get('id')
        if item_id:
            conn = get_db_connection()
            if target == 'student_stats':
                conn.execute('DELETE FROM student_stats WHERE id = ?', (item_id,))
                flash('បានលុបទិន្នន័យស្ថិតិសិស្សដោយជោគជ័យ!', 'success')
            elif target == 'exam_results':
                conn.execute('DELETE FROM exam_results WHERE id = ?', (item_id,))
                flash('បានលុបទិន្នន័យលទ្ធផលប្រឡងដោយជោគជ័យ!', 'success')
            conn.commit()
            conn.close()
        return redirect(url_for('students', level=level_filter))

    conn = get_db_connection()

    query = '''
        SELECT st.*, sc.name_kh as school_name, sc.level as school_level
        FROM student_stats st
        JOIN schools sc ON st.school_id = sc.id
    '''
    params = []
    if level_filter != 'all':
        query += ' WHERE sc.level = ?'
        params.append(level_filter)

    query += ' ORDER BY st.id DESC'
    stats_list = conn.execute(query, params).fetchall()

    if level_filter in ['community_preschool', 'state_preschool', 'primary']:
        exam_list = []
    elif level_filter == 'secondary':
        exam_list = conn.execute('''
            SELECT ex.*, sc.name_kh as school_name
            FROM exam_results ex
            JOIN schools sc ON ex.school_id = sc.id
            WHERE ex.exam_type = 'diploma_grade_9'
            ORDER BY ex.id DESC
        ''').fetchall()
    elif level_filter == 'high':
        exam_list = conn.execute('''
            SELECT ex.*, sc.name_kh as school_name
            FROM exam_results ex
            JOIN schools sc ON ex.school_id = sc.id
            WHERE ex.exam_type = 'bacii_grade_12'
            ORDER BY ex.id DESC
        ''').fetchall()
    else:
        exam_list = conn.execute('''
            SELECT ex.*, sc.name_kh as school_name
            FROM exam_results ex
            JOIN schools sc ON ex.school_id = sc.id
            ORDER BY ex.id DESC
        ''').fetchall()

    schools_list = conn.execute('SELECT id, name_kh, level FROM schools').fetchall()

    level_counts = {
        'all': conn.execute('SELECT COUNT(*) FROM student_stats').fetchone()[0],
        'community_preschool': conn.execute("SELECT COUNT(*) FROM student_stats st JOIN schools sc ON st.school_id = sc.id WHERE sc.level = 'community_preschool'").fetchone()[0],
        'state_preschool': conn.execute("SELECT COUNT(*) FROM student_stats st JOIN schools sc ON st.school_id = sc.id WHERE sc.level IN ('state_preschool', 'preschool')").fetchone()[0],
        'primary': conn.execute("SELECT COUNT(*) FROM student_stats st JOIN schools sc ON st.school_id = sc.id WHERE sc.level = 'primary'").fetchone()[0],
        'secondary': conn.execute("SELECT COUNT(*) FROM student_stats st JOIN schools sc ON st.school_id = sc.id WHERE sc.level = 'secondary'").fetchone()[0],
        'high': conn.execute("SELECT COUNT(*) FROM student_stats st JOIN schools sc ON st.school_id = sc.id WHERE sc.level = 'high'").fetchone()[0],
    }

    academic_years = [r[0] for r in conn.execute('SELECT DISTINCT academic_year FROM student_stats ORDER BY academic_year DESC').fetchall() if r[0]]

    conn.close()

    return render_template(
        'students.html',
        active_page='students',
        stats_list=stats_list,
        exam_list=exam_list,
        schools=schools_list,
        level_filter=level_filter,
        level_counts=level_counts,
        academic_years=academic_years
    )

@app.route('/students/add_stats', methods=['POST'])
@login_required
def add_student_stats():
    stat_id = request.form.get('stat_id')
    school_id = request.form.get('school_id')
    academic_year = request.form.get('academic_year')
    total_students = request.form.get('total_students', 0)
    female_students = request.form.get('female_students', 0)
    disabled_students = request.form.get('disabled_students', 0)
    disabled_female = request.form.get('disabled_female', 0)
    scholarship_students = request.form.get('scholarship_students', 0)
    scholarship_female = request.form.get('scholarship_female', 0)
    dropout_count = request.form.get('dropout_count', 0)
    dropout_female = request.form.get('dropout_female', 0)
    repetition_count = request.form.get('repetition_count', 0)

    conn = get_db_connection()
    if stat_id:
        conn.execute('''
            UPDATE student_stats SET school_id=?, academic_year=?, total_students=?, female_students=?, disabled_students=?, disabled_female=?, scholarship_students=?, scholarship_female=?, dropout_count=?, dropout_female=?, repetition_count=?
            WHERE id=?
        ''', (school_id, academic_year, total_students, female_students, disabled_students, disabled_female, scholarship_students, scholarship_female, dropout_count, dropout_female, repetition_count, stat_id))
        flash('បានកែប្រែស្ថិតិសិស្សដោយជោគជ័យ!', 'success')
    else:
        conn.execute('''
            INSERT INTO student_stats (school_id, academic_year, total_students, female_students, disabled_students, disabled_female, scholarship_students, scholarship_female, dropout_count, dropout_female, repetition_count)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (school_id, academic_year, total_students, female_students, disabled_students, disabled_female, scholarship_students, scholarship_female, dropout_count, dropout_female, repetition_count))
        flash('បានបញ្ជូនស្ថិតិសិស្សដោយជោគជ័យ!', 'success')
    conn.commit()
    conn.close()
    return redirect(url_for('students'))

@app.route('/students/add_exam', methods=['POST'])
@login_required
def add_exam_result():
    school_id = request.form.get('school_id')
    academic_year = request.form.get('academic_year')
    exam_type = request.form.get('exam_type')
    total_candidates = request.form.get('total_candidates', 0)
    female_candidates = request.form.get('female_candidates', 0)
    total_passed = request.form.get('total_passed', 0)
    female_passed = request.form.get('female_passed', 0)
    grade_a = request.form.get('grade_a', 0)
    grade_b = request.form.get('grade_b', 0)
    grade_c = request.form.get('grade_c', 0)
    grade_d = request.form.get('grade_d', 0)
    grade_e = request.form.get('grade_e', 0)

    conn = get_db_connection()
    conn.execute('''
        INSERT INTO exam_results (school_id, academic_year, exam_type, total_candidates, female_candidates, total_passed, female_passed, grade_a, grade_b, grade_c, grade_d, grade_e)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (school_id, academic_year, exam_type, total_candidates, female_candidates, total_passed, female_passed, grade_a, grade_b, grade_c, grade_d, grade_e))
    conn.commit()
    conn.close()
    flash('បានបញ្ជូនលទ្ធផលប្រឡងដោយជោគជ័យ!', 'success')
    return redirect(url_for('students'))

# --- MODULE 4: YOUTH & SPORTS MANAGEMENT ---
@app.route('/youth_sports')
@login_required
def youth_sports():
    conn = get_db_connection()
    athletes_list = conn.execute('''
        SELECT sa.*, sc.name_kh as school_name
        FROM sports_athletes sa
        JOIN schools sc ON sa.school_id = sc.id
        ORDER BY sa.id DESC
    ''').fetchall()

    youth_list = conn.execute('SELECT * FROM youth_clubs ORDER BY id DESC').fetchall()
    schools_list = conn.execute('SELECT id, name_kh FROM schools').fetchall()
    conn.close()

    return render_template('youth_sports.html', active_page='youth_sports', athletes_list=athletes_list, youth_list=youth_list, schools=schools_list)

@app.route('/youth_sports/add_athlete', methods=['POST'])
@login_required
def add_athlete():
    athlete_name = request.form.get('athlete_name')
    sex = request.form.get('sex')
    school_id = request.form.get('school_id')
    sport_type = request.form.get('sport_type')
    achievement = request.form.get('achievement')

    conn = get_db_connection()
    conn.execute('''
        INSERT INTO sports_athletes (event_id, school_id, athlete_name, sex, sport_type, achievement)
        VALUES (1, ?, ?, ?, ?, ?)
    ''', (school_id, athlete_name, sex, sport_type, achievement))
    conn.commit()
    conn.close()
    flash('បានបន្ថែមឈ្មោះកីឡាករជ័យលាភីដោយជោគជ័យ!', 'success')
    return redirect(url_for('youth_sports'))

@app.route('/youth_sports/add_club', methods=['POST'])
@login_required
def add_youth_club():
    title = request.form.get('title')
    club_type = request.form.get('club_type')
    location = request.form.get('location')
    leader_name = request.form.get('leader_name')
    total_members = request.form.get('total_members', 0)
    female_members = request.form.get('female_members', 0)
    activity_date = request.form.get('activity_date')
    description = request.form.get('description')

    conn = get_db_connection()
    conn.execute('''
        INSERT INTO youth_clubs (title, club_type, location, leader_name, total_members, female_members, activity_date, description)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    ''', (title, club_type, location, leader_name, total_members, female_members, activity_date, description))
    conn.commit()
    conn.close()
    flash('បានបន្ថែមក្លឹប/សកម្មភាពយុវជនដោយជោគជ័យ!', 'success')
    return redirect(url_for('youth_sports'))

# --- MODULE 5: ADMINISTRATION & DOCUMENTS ---
@app.route('/documents')
@login_required
def documents():
    sub = request.args.get('sub')
    action = request.args.get('action')
    if action == 'delete':
        target = request.args.get('target')
        item_id = request.args.get('id')
        if item_id:
            conn = get_db_connection()
            if target == 'document':
                conn.execute('DELETE FROM documents WHERE id = ?', (item_id,))
                flash('បានលុបលិខិតផ្លូវការដោយជោគជ័យ!', 'success')
                sub = sub or 'letters'
            elif target == 'mission':
                conn.execute('DELETE FROM mission_orders WHERE id = ?', (item_id,))
                flash('បានលុបលិខិតបង្គាប់ការដោយជោគជ័យ!', 'success')
                sub = sub or 'missions'
            conn.commit()
            conn.close()
        if sub:
            return redirect(url_for('documents', sub=sub))
        return redirect(url_for('documents'))

    conn = get_db_connection()
    docs_list = conn.execute('SELECT * FROM documents ORDER BY id DESC').fetchall()
    missions_list = conn.execute('SELECT * FROM mission_orders ORDER BY id DESC').fetchall()
    conn.close()
    return render_template('documents.html', active_page='documents', docs_list=docs_list, missions_list=missions_list)

@app.route('/documents/add', methods=['POST'])
@login_required
def add_document():
    doc_id = request.form.get('doc_id')
    doc_number = request.form.get('doc_number')
    doc_type = request.form.get('doc_type')
    title = request.form.get('title')
    source_destination = request.form.get('source_destination')
    doc_date = request.form.get('doc_date')
    category = request.form.get('category', 'លិខិតផ្លូវការ')
    scanned_file = request.form.get('scanned_file')

    conn = get_db_connection()
    if doc_id:
        conn.execute('''
            UPDATE documents SET doc_number=?, doc_type=?, title=?, source_destination=?, doc_date=?, category=?
            WHERE id=?
        ''', (doc_number, doc_type, title, source_destination, doc_date, category, doc_id))
        flash('បានធ្វើបច្ចុប្បន្នភាពព័ត៌មានលិខិតផ្លូវការដោយជោគជ័យ!', 'success')
    else:
        conn.execute('''
            INSERT INTO documents (doc_number, doc_type, title, source_destination, doc_date, category, scanned_file)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (doc_number, doc_type, title, source_destination, doc_date, category, scanned_file))
        flash('បានកត់ត្រាលិខិតផ្លូវការដោយជោគជ័យ!', 'success')
    conn.commit()
    conn.close()
    return redirect(url_for('documents', sub='letters'))

@app.route('/documents/add_mission', methods=['POST'])
@login_required
@role_required('admin', 'chief')
def add_mission_order():
    mission_id = request.form.get('mission_id')
    mission_code = request.form.get('mission_code')
    title = request.form.get('title')
    officer_names = request.form.get('officer_names')
    destination_schools = request.form.get('destination_schools')
    start_date = request.form.get('start_date')
    end_date = request.form.get('end_date')
    purpose = request.form.get('purpose', '')
    ref_doc = request.form.get('ref_doc', '')
    lunar_date_str = request.form.get('lunar_date_str', '')
    mission_type = request.form.get('mission_type', 'local')

    try:
        days_count = int(request.form.get('days_count') or 1)
    except (ValueError, TypeError):
        days_count = 1
    try:
        nights_count = int(request.form.get('nights_count') or 0)
    except (ValueError, TypeError):
        nights_count = 0
    def clean_float(val, default=0.0):
        if not val:
            return default
        try:
            return float(str(val).replace(',', '').replace(' ', '').strip())
        except (ValueError, TypeError):
            return default

    travel_cost = clean_float(request.form.get('travel_cost'), 0.0)
    pocket_rate = clean_float(request.form.get('pocket_rate'), 0.0)
    pocket_total = clean_float(request.form.get('pocket_total'), pocket_rate * days_count)
    food_rate = clean_float(request.form.get('food_rate'), 0.0)
    food_total = clean_float(request.form.get('food_total'), food_rate * days_count)
    hotel_rate = clean_float(request.form.get('hotel_rate'), 0.0)
    hotel_total = clean_float(request.form.get('hotel_total'), hotel_rate * nights_count)
    total_cost = clean_float(request.form.get('total_cost'), travel_cost + pocket_total + food_total + hotel_total)

    conn = get_db_connection()
    if mission_id:
        conn.execute('''
            UPDATE mission_orders SET
                mission_code=?, title=?, officer_names=?, destination_schools=?,
                start_date=?, end_date=?, purpose=?, ref_doc=?, lunar_date_str=?,
                mission_type=?, days_count=?, nights_count=?, travel_cost=?,
                pocket_rate=?, pocket_total=?, food_rate=?, food_total=?,
                hotel_rate=?, hotel_total=?, total_cost=?
            WHERE id=?
        ''', (
            mission_code, title, officer_names, destination_schools,
            start_date, end_date, purpose, ref_doc, lunar_date_str,
            mission_type, days_count, nights_count, travel_cost,
            pocket_rate, pocket_total, food_rate, food_total,
            hotel_rate, hotel_total, total_cost, mission_id
        ))
        flash('បានធ្វើបច្ចុប្បន្នភាពលិខិតបញ្ជាបេសកកម្មដោយជោគជ័យ!', 'success')
    else:
        conn.execute('''
            INSERT INTO mission_orders (
                mission_code, title, officer_names, destination_schools,
                start_date, end_date, purpose, status, ref_doc, lunar_date_str,
                mission_type, days_count, nights_count, travel_cost,
                pocket_rate, pocket_total, food_rate, food_total,
                hotel_rate, hotel_total, total_cost
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, 'approved', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            mission_code, title, officer_names, destination_schools,
            start_date, end_date, purpose, ref_doc, lunar_date_str,
            mission_type, days_count, nights_count, travel_cost,
            pocket_rate, pocket_total, food_rate, food_total,
            hotel_rate, hotel_total, total_cost
        ))
        if mission_type == 'cross_district':
            flash('បានបង្កើតបេសកកម្មឆ្លងស្រុកដោយជោគជ័យ!', 'success')
        else:
            flash('បានចេញលិខិតបញ្ជាបេសកកម្មក្នុងស្រុកដោយជោគជ័យ!', 'success')
    conn.commit()
    conn.close()
    return redirect(url_for('documents', sub='missions'))

def khmer_num_to_arabic(str_val):
    if not str_val:
        return ''
    khmer_digits = {'០':'0', '១':'1', '២':'2', '៣':'3', '៤':'4', '៥':'5', '៦':'6', '៧':'7', '៨':'8', '៩':'9'}
    return ''.join(khmer_digits.get(c, c) for c in str(str_val))

@app.route('/api/parse_mission_file', methods=['POST'])
@login_required
def parse_mission_file():
    try:
        file = request.files.get('file')
        client_text = request.form.get('client_text', '')

        extracted_text = ""

        if client_text and len(client_text.strip()) > 5:
            extracted_text = client_text
        elif file:
            filename = file.filename.lower()
            try:
                if filename.endswith('.pdf'):
                    import pypdf
                    reader = pypdf.PdfReader(file.stream)
                    for page in reader.pages:
                        t = page.extract_text()
                        if t:
                            extracted_text += t + "\n"
                else:
                    raw = file.stream.read().decode('utf-8', errors='ignore')
                    if not raw.startswith('%PDF-') and 'endobj' not in raw:
                        extracted_text = raw
            except Exception:
                extracted_text = ""

        # Clean binary PDF noise lines
        clean_lines = []
        for line in extracted_text.splitlines():
            line_str = line.strip()
            if any(token in line_str for token in ['%PDF-', '/CreationDate', 'AppleWebKit', '/Type', '/Font', 'endobj', 'stream', '/MediaBox', '/Producer']):
                continue
            clean_lines.append(line_str)
        extracted_text = "\n".join(clean_lines)

        parsed_data = {
            'mission_code': '',
            'lunar_date_str': '',
            'ref_doc': '',
            'title': '',
            'destination_schools': '',
            'start_date': '',
            'end_date': '',
            'officer_names': ''
        }

        if extracted_text and len(extracted_text.strip()) > 3:
            # Code match
            code_match = re.search(r'លេខ[៖:]?\s*([០-៩0-9/-]+)', extracted_text)
            if code_match:
                parsed_data['mission_code'] = code_match.group(1).strip()

            # Lunar date match
            lunar_match = re.search(r'(ថ្ងៃ\s+[^\n\r]+(?:ឆ្នាំ[^\n\r]+)?)', extracted_text)
            if lunar_match:
                parsed_data['lunar_date_str'] = lunar_match.group(1).strip()

            # Ref doc match
            ref_match = re.search(r'យោង[៖:]?\s*([^\n\r]+)', extracted_text)
            if ref_match:
                parsed_data['ref_doc'] = ref_match.group(1).strip()

            # Title match
            title_match = re.search(r'កម្មវត្ថុ[៖:]?\s*([^\n\r]+)', extracted_text)
            if title_match:
                parsed_data['title'] = title_match.group(1).strip()

            # Destination match
            dest_match = re.search(r'(?:ទីកន្លែង|នៅ)[៖:]?\s*([^\n\r]+)', extracted_text)
            if dest_match:
                parsed_data['destination_schools'] = dest_match.group(1).strip()

            # Officers list (strictly matching numbered lines with Khmer or Arabic digits)
            officer_matches = re.findall(r'([១២៣៤៥៦៧៨៩០1-9][\.\)]\s*[^\n\r]+)', extracted_text)
            if officer_matches:
                valid_officers = [o.strip() for o in officer_matches if not re.search(r'(Windows|Safari|Mozilla|CreationDate|AppleWebKit)', o)]
                if valid_officers:
                    parsed_data['officer_names'] = "\n".join(valid_officers)

            # Date Range Extraction
            khmer_months = {
                'មករា': '01', 'កុម្ភៈ': '02', 'មីនា': '03', 'មេសា': '04',
                'ឧសភា': '05', 'មិថុនា': '06', 'កក្កដា': '07', 'សីហា': '08',
                'កញ្ញា': '09', 'តុលា': '10', 'វិច្ឆិកា': '11', 'ធ្នូ': '12'
            }
            date_match = re.search(r'ចាប់ពីថ្ងៃទី\s*([០-៩0-9]+)\s*(?:ដល់|ទៅ|ដល់ថ្ងៃទី)\s*([០-៩0-9]+)\s*ខែ\s*([^\s]+)\s*ឆ្នាំ\s*([០-៩0-9]{4})', extracted_text)
            if date_match:
                start_day = khmer_num_to_arabic(date_match.group(1)).zfill(2)
                end_day = khmer_num_to_arabic(date_match.group(2)).zfill(2)
                month_name = date_match.group(3).strip()
                month_code = khmer_months.get(month_name, '01')
                year_num = khmer_num_to_arabic(date_match.group(4))

                parsed_data['start_date'] = f"{year_num}-{month_code}-{start_day}"
                parsed_data['end_date'] = f"{year_num}-{month_code}-{end_day}"

        # Only apply generic sample defaults IF no text could be extracted at all from the user's original document
        has_real_text = bool(extracted_text and len(extracted_text.strip()) > 10)
        if not has_real_text:
            import datetime
            today = datetime.date.today()

            if not parsed_data['mission_code']:
                fn_nums = re.findall(r'\d+', file.filename if file else '')
                parsed_data['mission_code'] = fn_nums[0] if fn_nums else '១៥៨'

            if not parsed_data['lunar_date_str']:
                parsed_data['lunar_date_str'] = 'ថ្ងៃ ច័ន្ទ ០៥ កើត ខែផល្គុន ឆ្នាំថោះ ព.ស ២៥៦៧'

            if not parsed_data['ref_doc']:
                parsed_data['ref_doc'] = 'លិខិតអញ្ជើញលេខ ១៥៨ អយក.រប ចុះថ្ងៃទី១៥ ខែសីហា ឆ្នាំ២០២៦ របស់មន្ទីរអប់រំ យុវជន និងកីឡាខេត្តសៀមរាប'

            if not parsed_data['title']:
                parsed_data['title'] = 'ចូលរួមបេសកកម្មត្រួតពិនិត្យ វាយតម្លៃ និងពង្រឹងគុណភាពការងារអប់រំក្នុងស្រុកអង្គរធំ'

            if not parsed_data['destination_schools']:
                parsed_data['destination_schools'] = 'សាលារៀន និងគ្រឹះស្ថានសិក្សាក្នុងស្រុកអង្គរធំ ខេត្តសៀមរាប'

            if not parsed_data['start_date']:
                parsed_data['start_date'] = today.strftime('%Y-%m-%d')
                parsed_data['end_date'] = (today + datetime.timedelta(days=3)).strftime('%Y-%m-%d')

            if not parsed_data['officer_names']:
                parsed_data['officer_names'] = "១. លោកស្រី គិម ផាន់  នាយិកាវិទ្យាល័យតេជោហ៊ុនសែនអង្គរធំ\n២. លោក ជា តិចស៊ីន  នាយកអនុវិទ្យាល័យអង្គរធំ"

        return jsonify({
            'success': True,
            'parsed_data': parsed_data,
            'filename': file.filename if file else 'document',
            'has_real_text': has_real_text,
            'raw_text_length': len(extracted_text.strip())
        })

    except Exception as err:
        return jsonify({
            'success': False,
            'message': f'ការអានឯកសារមានបញ្ហា៖ {str(err)}'
        }), 200

def number_to_khmer_words(number):
    if not number:
        return 'សូន្យរៀលគត់'
    digits = ['សូន្យ', 'មួយ', 'ពីរ', 'បី', 'បួន', 'ប្រាំ', 'ប្រាំមួយ', 'ប្រាំពីរ', 'ប្រាំបី', 'ប្រាំបួន']
    tens = ['', 'ដប់', 'ម្ភៃ', 'សាមសិប', 'សែសិប', 'ហាសិប', 'ហុកសិប', 'ចិបសិប', 'ប៉ែតសិប', 'កៅសិប']
    n = int(number)
    if n == 0:
        return 'សូន្យរៀលគត់'

    def convert_below_thousand(val):
        res = ''
        if val >= 100:
            res += digits[val // 100] + 'រយ'
            val %= 100
        if val >= 10:
            res += tens[val // 10]
            val %= 10
        if val > 0:
            res += digits[val]
        return res

    res = ''
    if n >= 1000000:
        res += convert_below_thousand(n // 1000000) + 'លាន '
        n %= 1000000
    if n >= 1000:
        res += convert_below_thousand(n // 1000) + 'ពាន់ '
        n %= 1000
    if n > 0:
        res += convert_below_thousand(n)
    return res.strip() + 'រៀលគត់'

def get_officer_summary(missions):
    import re
    officer_map = {}
    conn = get_db_connection()
    office_staff_rows = conn.execute('SELECT staff_id_num, name_kh, name_en, sex, position, bank_account FROM office_staff').fetchall()
    staff_rows = conn.execute('SELECT staff_id_num, name_kh, name_en, sex, position, bank_account FROM staff').fetchall()
    conn.close()

    staff_meta = {}

    def add_to_meta(row):
        name_raw = (row['name_kh'] or '').strip()
        c_name = re.sub(r'^[០-៩0-9\.\s\-—]+', '', name_raw).strip()
        c_name = c_name.replace('លោកស្រី', '').replace('លោក', '').strip()
        data = {
            'staff_id_num': row.get('staff_id_num', ''),
            'name_en': row.get('name_en', ''),
            'sex': 'ស្រី' if row.get('sex') == 'F' else 'ប្រុស',
            'position': row.get('position', 'មន្ត្រី'),
            'bank_account': row.get('bank_account', '')
        }
        if name_raw:
            staff_meta[name_raw] = data
        if c_name:
            staff_meta[c_name] = data

    # 1. School staff
    for s in staff_rows:
        add_to_meta(s)

    # 2. Office staff (takes highest priority for officer Latin names & details)
    for os in office_staff_rows:
        add_to_meta(os)

    default_accounts = {
        'ស្រី យិន': {'id': '1721700234', 'en': 'Srey Yan', 'account': '3451-01-636512-1-8'},
        'ហួត ស': {'id': '1671700093', 'en': 'Huot Sor', 'account': '3451-01-636478-1-3'},
        'ស្ងៀម សារឿង': {'id': '2831700060', 'en': 'Siem Salo', 'account': '3451-01-636571-1-8'},
        'ស្ងៀម សារឿ': {'id': '2831700060', 'en': 'Siem Salo', 'account': '3451-01-636571-1-8'},
        'ឡី លាត់': {'id': '1871700111', 'en': 'Ley Loat', 'account': '3451-01-646849-1-2'},
        'ចក់ សុភាព': {'id': '2951700023', 'en': 'Chork Sopheap', 'account': '3451-01-957506-1-9'},
        'ហាស សំបូរ': {'id': '1871700110', 'en': 'Has Sambo', 'account': '3451-01-468703-2-1'},
        'សែន ហួត': {'id': '2941700070', 'en': 'Sen Houch', 'account': '3451-01-957879-1-1'},
    }

    for m in missions:
        raw_names = m.get('officer_names', '')
        name_list = [n.strip() for n in raw_names.replace(',', '\n').split('\n') if n.strip()]
        if not name_list:
            name_list = [raw_names.strip()]

        allowance = float(m.get('allowance_amount') or 40000.0)

        for name in name_list:
            clean_name = re.sub(r'^[០-៩0-9\.\s\-—]+', '', name).strip()
            clean_name = clean_name.replace('លោកស្រី', '').replace('លោក', '').strip()
            if not clean_name:
                continue

            if clean_name not in officer_map:
                meta = staff_meta.get(clean_name, {})
                def_data = default_accounts.get(clean_name, {})

                staff_id_num = meta.get('staff_id_num') or def_data.get('id') or '1721700123'
                name_en = meta.get('name_en') or def_data.get('en') or clean_name
                bank_account = meta.get('bank_account') or def_data.get('account') or '3451-01-636000-1-0'

                officer_map[clean_name] = {
                    'name': clean_name,
                    'staff_id_num': staff_id_num,
                    'name_en': name_en,
                    'bank_account': bank_account,
                    'sex': meta.get('sex', 'ប្រុស'),
                    'position': meta.get('position', 'មន្ត្រី'),
                    'trip_count': 0,
                    'rate_per_trip': allowance,
                    'total_amount': 0.0
                }

            officer_map[clean_name]['trip_count'] += 1
            officer_map[clean_name]['total_amount'] += allowance

    summary_list = list(officer_map.values())
    total_summary_trips = sum(o['trip_count'] for o in summary_list)
    total_summary_amount = sum(o['total_amount'] for o in summary_list)

    return summary_list, total_summary_trips, total_summary_amount

def format_khmer_date(date_str):
    if not date_str:
        return ''
    parts = date_str.split('-')
    if len(parts) == 3:
        y, m, d = parts
        return f"{d}-{m}-{y}"
    return date_str

def group_missions_by_period(missions):
    import re
    months_kh = {
        '01': 'មករា', '02': 'កុម្ភៈ', '03': 'មីនា', '04': 'មេសា',
        '05': 'ឧសភា', '06': 'មិថុនា', '07': 'កក្កដា', '08': 'សីហា',
        '09': 'កញ្ញា', '10': 'តុលា', '11': 'វិច្ឆិកា', '12': 'ធ្នូ'
    }

    def to_kh_date_label(d_str):
        if not d_str or len(d_str.split('-')) != 3:
            return d_str
        y, m, d = d_str.split('-')
        m_name = months_kh.get(m, m)
        return f"{int(d):02d} ខែ {m_name} ឆ្នាំ{y}"

    groups_map = {}
    sorted_missions = sorted(missions, key=lambda m: (m.get('start_date') or '', m.get('id') or 0))

    for m in sorted_missions:
        s_date = m.get('start_date') or ''
        parts = s_date.split('-')
        if len(parts) == 3:
            group_key = f"{parts[0]}-{parts[1]}"
        else:
            group_key = "other"
        if group_key not in groups_map:
            groups_map[group_key] = []
        groups_map[group_key].append(m)

    result_groups = []
    global_counter = 1

    for g_key in sorted(groups_map.keys()):
        m_list = groups_map[g_key]
        if not m_list:
            continue

        min_start = min(m.get('start_date') or '' for m in m_list)
        max_end = max((m.get('end_date') or m.get('start_date') or '') for m in m_list)

        group_title = f"ចាប់ពីថ្ងៃទី {to_kh_date_label(min_start)} ដល់ថ្ងៃទី {to_kh_date_label(max_end)}"
        group_items = []
        group_subtotal = 0

        for m in m_list:
            raw_officers = m.get('officer_names') or ''
            lines = [l.strip() for l in raw_officers.replace(',', '\n').split('\n') if l.strip()]
            if not lines:
                lines = [raw_officers]

            mission_allowance = m.get('allowance_amount') or 40000.0

            m_item = {
                'id': m.get('id'),
                'mission_code': m.get('mission_code'),
                'title': m.get('title'),
                'destination_schools': m.get('destination_schools'),
                'start_date': format_khmer_date(m.get('start_date')),
                'end_date': format_khmer_date(m.get('end_date')),
                'remarks': m.get('remarks') or '',
                'officers': [],
                'rowspan': len(lines)
            }

            for idx, off_name in enumerate(lines):
                clean_name = re.sub(r'^[០-៩0-9\.\s\-—]+', '', off_name).strip()
                if not clean_name:
                    clean_name = off_name

                m_item['officers'].append({
                    'index': global_counter,
                    'name': clean_name,
                    'amount': mission_allowance,
                    'is_first': (idx == 0)
                })
                group_subtotal += mission_allowance
                global_counter += 1

            group_items.append(m_item)

        result_groups.append({
            'group_title': group_title,
            'subtotal': group_subtotal,
            'missions': group_items
        })

    return result_groups

@app.route('/documents/mission_report')
@app.route('/reports/mission_expenses')
@login_required
def mission_expense_report():
    start_date = request.args.get('start_date', '')
    end_date = request.args.get('end_date', '')

    conn = get_db_connection()
    query = 'SELECT * FROM mission_orders WHERE 1=1'
    params = []

    if start_date:
        query += ' AND start_date >= ?'
        params.append(start_date)
    if end_date:
        query += ' AND end_date <= ?'
        params.append(end_date)

    query += ' ORDER BY start_date ASC, id ASC'
    missions = conn.execute(query, params).fetchall()
    conn.close()

    mission_groups = group_missions_by_period(missions)
    total_allowance = sum(g['subtotal'] for g in mission_groups)
    total_allowance_words = number_to_khmer_words(total_allowance)

    officer_summary, total_summary_trips, total_summary_amount = get_officer_summary(missions)
    total_summary_amount_words = number_to_khmer_words(total_summary_amount)

    start_date_display = start_date if start_date else '១៧/០៤/២០២៦'
    end_date_display = end_date if end_date else '០៩/០៦/២០២៦'

    return render_template(
        'mission_expense_report.html',
        active_page='documents',
        missions=missions,
        mission_groups=mission_groups,
        total_allowance=total_allowance,
        total_allowance_words=total_allowance_words,
        officer_summary=officer_summary,
        total_summary_trips=total_summary_trips,
        total_summary_amount=total_summary_amount,
        total_summary_amount_words=total_summary_amount_words,
        start_date=start_date,
        end_date=end_date,
        start_date_display=start_date_display,
        end_date_display=end_date_display
    )

# --- MODULE 6: BUDGET & ASSETS MANAGEMENT ---
@app.route('/budget_assets')
@login_required
def budget_assets():
    conn = get_db_connection()
    budget_list = conn.execute('''
        SELECT b.*, sc.name_kh as school_name
        FROM budgets b
        JOIN schools sc ON b.school_id = sc.id
        ORDER BY b.id DESC
    ''').fetchall()

    assets_list = conn.execute('''
        SELECT a.*, sc.name_kh as school_name
        FROM assets a
        JOIN schools sc ON a.school_id = sc.id
        ORDER BY a.id DESC
    ''').fetchall()

    schools_list = conn.execute('SELECT id, name_kh FROM schools').fetchall()
    academic_years = [r[0] for r in conn.execute('SELECT DISTINCT academic_year FROM student_stats ORDER BY academic_year DESC').fetchall() if r[0]]
    conn.close()

    return render_template('budget_assets.html', active_page='budget_assets', budget_list=budget_list, assets_list=assets_list, schools=schools_list, academic_years=academic_years)

@app.route('/budget_assets/add_budget', methods=['POST'])
@login_required
@role_required('admin', 'chief')
def add_budget():
    school_id = request.form.get('school_id')
    academic_year = request.form.get('academic_year')
    budget_type = request.form.get('budget_type')
    allocated_amount = request.form.get('allocated_amount', 0)
    spent_amount = request.form.get('spent_amount', 0)
    notes = request.form.get('notes')

    conn = get_db_connection()
    conn.execute('''
        INSERT INTO budgets (school_id, academic_year, budget_type, allocated_amount, spent_amount, notes)
        VALUES (?, ?, ?, ?, ?, ?)
    ''', (school_id, academic_year, budget_type, allocated_amount, spent_amount, notes))
    conn.commit()
    conn.close()
    flash('បានកត់ត្រាកញ្ចប់ថវិកាដោយជោគជ័យ!', 'success')
    return redirect(url_for('budget_assets'))

@app.route('/budget_assets/add_asset', methods=['POST'])
@login_required
def add_asset():
    school_id = request.form.get('school_id')
    asset_name = request.form.get('asset_name')
    category = request.form.get('category')
    quantity = request.form.get('quantity', 0)
    condition = request.form.get('condition', 'good')
    date_allocated = request.form.get('date_allocated')

    conn = get_db_connection()
    conn.execute('''
        INSERT INTO assets (school_id, asset_name, category, quantity, condition, date_allocated)
        VALUES (?, ?, ?, ?, ?, ?)
    ''', (school_id, asset_name, category, quantity, condition, date_allocated))
    conn.commit()
    conn.close()
    flash('បានបន្ថែមទិន្នន័យសម្ភារ/ទ្រព្យសម្បត្តិដោយជោគជ័យ!', 'success')
    return redirect(url_for('budget_assets'))

# --- MODULE 7: REPORTS & EXPORTS ---
@app.route('/reports')
@login_required
def reports():
    return render_template('reports.html', active_page='reports')

@app.route('/exports/excel/<report_type>')
@login_required
def export_excel(report_type):
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Report"

        header_font = Font(name='Segoe UI', size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0F2B5C", end_color="0F2B5C", fill_type="solid")

        conn = get_db_connection()
        if report_type == 'schools_staff':
            ws.append(["កូដសាលា", "ឈ្មោះសាលារៀន", "កម្រិត", "ឃុំ/ភូមិ", "នាយកសាលា", "លេខទូរស័ព្ទ"])
            rows = conn.execute('SELECT code, name_kh, level, commune || " / " || village, principal_name, principal_phone FROM schools').fetchall()
            for r in rows:
                ws.append(list(r))
        elif report_type == 'students_exam':
            ws.append(["សាលារៀន", "ឆ្នាំសិក្សា", "ប្រភេទប្រឡង", "បេក្ខជនសរុប", "ប្រឡងជាប់", "និទ្ទេស A", "និទ្ទេស B"])
            rows = conn.execute('''
                SELECT sc.name_kh, ex.academic_year, ex.exam_type, ex.total_candidates, ex.total_passed, ex.grade_a, ex.grade_b
                FROM exam_results ex JOIN schools sc ON ex.school_id = sc.id
            ''').fetchall()
            for r in rows:
                ws.append(list(r.values()) if hasattr(r, 'values') else list(r))
        elif report_type == 'mission_expenses':
            ws.append(["ល.រ", "គោត្តនាម នាម", "លេខលិខិត", "កាលបរិច្ឆេទលិខិត", "កម្មវត្ថុនៃការចុះបេសកកម្មបំរើសកម្មភាព", "ទីកន្លែង", "ថ្ងៃចាប់ផ្តើម", "ថ្ងៃត្រឡប់", "ប្រាក់ឧបត្ថម្ភ (៛)", "ផ្សេងៗ"])
            rows = conn.execute('SELECT officer_names, mission_code, start_date, title, destination_schools, start_date, end_date, COALESCE(allowance_amount, 40000), COALESCE(remarks, "") FROM mission_orders ORDER BY start_date ASC').fetchall()
            for idx, r in enumerate(rows, 1):
                r_values = list(r.values()) if hasattr(r, 'values') else list(r)
                ws.append([idx] + r_values)
        elif report_type == 'officer_summary':
            ws.append(["ល.រ", "គោត្តនាម និង នាម", "ភេទ", "តួនាទី", "ចំនួនបានចុះ (លើក)", "ចំនួនថវិកា ១ លើក (៛)", "ចំនួនថវិកាសរុប (៛)"])
            missions = conn.execute('SELECT * FROM mission_orders ORDER BY start_date ASC').fetchall()
            summary_list, _, _ = get_officer_summary(missions)
            for idx, o in enumerate(summary_list, 1):
                ws.append([idx, o['name'], o['sex'], o['position'], o['trip_count'], o['rate_per_trip'], o['total_amount']])
        else:
            ws.append(["សាលារៀន", "ឆ្នាំសិក្សា", "ប្រភេទថវិកា", "ថវិកាទទួលបាន", "ថវិកាចាយរួច"])
            rows = conn.execute('''
                SELECT sc.name_kh, b.academic_year, b.budget_type, b.allocated_amount, b.spent_amount
                FROM budgets b JOIN schools sc ON b.school_id = sc.id
            ''').fetchall()
            for r in rows:
                ws.append(list(r.values()) if hasattr(r, 'values') else list(r))

        conn.close()

        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill

        file_path = os.path.join(EXPORTS_DIR, f"{report_type}_report.xlsx")
        wb.save(file_path)
        return send_file(file_path, as_attachment=True, download_name=f"Angkor_Thom_{report_type}_report.xlsx")
    except Exception as e:
        flash(f'មិនអាច Export Excel បានទេ: {str(e)}', 'danger')
        return redirect(url_for('reports'))

# --- PROFILE ROUTE ---
@app.route('/profile', methods=['GET', 'POST'])
@login_required
def profile():
    conn = get_db_connection()
    user_id = session.get('user_id')

    if request.method == 'POST':
        full_name = request.form.get('full_name')
        phone = request.form.get('phone')
        position = request.form.get('position')
        new_password = request.form.get('new_password')

        if new_password and len(new_password.strip()) > 0:
            from database import hash_password
            pwd_hash = hash_password(new_password.strip())
            conn.execute('''
                UPDATE users SET full_name=?, phone=?, position=?, password_hash=?
                WHERE id=?
            ''', (full_name, phone, position, pwd_hash, user_id))
        else:
            conn.execute('''
                UPDATE users SET full_name=?, phone=?, position=?
                WHERE id=?
            ''', (full_name, phone, position, user_id))

        conn.commit()
        session['user_name'] = full_name
        flash('បានធ្វើបច្ចុប្បន្នភាពព័ត៌មានគណនីរបស់ខ្ញុំដោយជោគជ័យ!', 'success')
        return redirect(url_for('profile'))

    user_info = conn.execute('SELECT * FROM users WHERE id = ?', (user_id,)).fetchone()
    conn.close()
    return render_template('profile.html', active_page='profile', user_info=user_info)

@app.route('/exports/pdf/<report_type>')
@login_required
def export_pdf(report_type):
    conn = get_db_connection()
    schools_list = conn.execute('SELECT * FROM schools').fetchall()
    conn.close()

    html = f"""
    <!DOCTYPE html>
    <html lang="km">
    <head>
        <meta charset="UTF-8">
        <title>របាយការណ៍ផ្លូវការ - ស្រុកអង្គរធំ</title>
        <style>
            body {{ font-family: 'Kantumruy Pro', sans-serif; padding: 40px; color: #000; }}
            h2 {{ text-align: center; margin-bottom: 5px; }}
            p {{ text-align: center; color: #555; margin-top: 0; }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 20px; font-size: 13px; }}
            th, td {{ border: 1px solid #333; padding: 8px; text-align: left; }}
            th {{ background: #f0f4f8; }}
        </style>
    </head>
    <body onload="window.print()">
        <h2>ការិយាល័យអប់រំ យុវជន និងកីឡា ស្រុកអង្គរធំ</h2>
        <p>របាយការណ៍ផ្លូវការ {report_type.upper()}</p>
        <table>
            <thead>
                <tr><th>ល.រ</th><th>ឈ្មោះសាលារៀន</th><th>ទីតាំង & នាយកសាលា</th></tr>
            </thead>
            <tbody>
    """
    for idx, s in enumerate(schools_list, 1):
        html += f"<tr><td>{idx}</td><td>{s['name_kh']}</td><td>ឃុំ{s['commune']} - នាយក {s['principal_name']} ({s['principal_phone']})</td></tr>"

    html += """
            </tbody>
        </table>
    </body>
    </html>
    """
    return html

# --- MODULE 8: USER AUTHORIZATION ---
@app.route('/users')
@login_required
@role_required('admin')
def users():
    action = request.args.get('action')
    user_id = request.args.get('id')

    if action == 'delete' and user_id:
        conn = get_db_connection()
        if int(user_id) == session.get('user_id'):
            flash('លោកអ្នកមិនអាចលុបគណនីដែលកំពុងចូលប្រើប្រាស់បច្ចុប្បន្នបានទេ!', 'danger')
        else:
            conn.execute('DELETE FROM users WHERE id = ?', (user_id,))
            conn.commit()
            flash('បានលុបគណនីអ្នកប្រើប្រាស់ដោយជោគជ័យ!', 'success')
        conn.close()
        return redirect(url_for('users'))

    conn = get_db_connection()
    users_list = conn.execute('''
        SELECT u.*, sc.name_kh as school_name
        FROM users u
        LEFT JOIN schools sc ON u.school_id = sc.id
        ORDER BY u.id ASC
    ''').fetchall()
    schools_list = conn.execute('SELECT id, name_kh FROM schools').fetchall()
    conn.close()

    return render_template('users.html', active_page='users', users_list=users_list, schools=schools_list)

@app.route('/users/add', methods=['POST'])
@login_required
@role_required('admin')
def add_user():
    user_id = request.form.get('user_id')
    username = request.form.get('username')
    password = request.form.get('password')
    full_name = request.form.get('full_name')
    role = request.form.get('role')
    school_id = request.form.get('school_id') or None
    position = request.form.get('position')
    phone = request.form.get('phone')

    conn = get_db_connection()
    try:
        if user_id:
            if password and password.strip():
                pwd_hash = generate_password_hash(password)
                conn.execute('''
                    UPDATE users SET username=?, password_hash=?, full_name=?, role=?, school_id=?, position=?, phone=?
                    WHERE id=?
                ''', (username, pwd_hash, full_name, role, school_id, position, phone, user_id))
            else:
                conn.execute('''
                    UPDATE users SET username=?, full_name=?, role=?, school_id=?, position=?, phone=?
                    WHERE id=?
                ''', (username, full_name, role, school_id, position, phone, user_id))
            flash('បានកែប្រែគណនីអ្នកប្រើប្រាស់ដោយជោគជ័យ!', 'success')
        else:
            pwd_hash = generate_password_hash(password if password else '123456')
            conn.execute('''
                INSERT INTO users (username, password_hash, full_name, role, school_id, position, phone)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            ''', (username, pwd_hash, full_name, role, school_id, position, phone))
            flash('បានបង្កើតគណនីអ្នកប្រើប្រាស់ថ្មីដោយជោគជ័យ!', 'success')
        conn.commit()
    except sqlite3.IntegrityError:
        flash('ឈ្មោះគណនីនេះមានរួចហើយ!', 'danger')
    finally:
        conn.close()

    return redirect(url_for('users'))

# --- MODULE 9: AUDIT LOGS ---
@app.route('/logs')
@login_required
@role_required('admin')
def logs():
    conn = get_db_connection()
    logs_list = conn.execute('SELECT * FROM audit_logs ORDER BY id DESC LIMIT 500').fetchall()
    conn.close()
    return render_template('logs.html', active_page='logs', logs_list=logs_list)

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print(f"Starting Angkor Thom District Education Management System Server on port {port}...")
    app.run(host='0.0.0.0', port=port, debug=True)

